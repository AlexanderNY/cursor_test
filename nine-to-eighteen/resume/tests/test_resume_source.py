"""Unit tests for resume source parse + evidenced path skills."""
from __future__ import annotations

import importlib.util
import io
import sys
from pathlib import Path
from types import ModuleType

API_DIR = Path(__file__).resolve().parents[1]


def _load(name: str, relative: str) -> ModuleType:
    path = API_DIR / relative
    spec = importlib.util.spec_from_file_location(name, path)
    assert spec and spec.loader
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


# text_sanitize first (dependency of resume_source)
_sanitize = _load("text_sanitize", "services/text_sanitize.py")
sys.modules["services.text_sanitize"] = _sanitize

_skills = _load("resume_skills", "services/resume_skills.py")
sys.modules["services.resume_skills"] = _skills
sys.modules["resume_skills"] = _skills

services_pkg = ModuleType("services")
services_pkg.resume_skills = _skills
services_pkg.text_sanitize = _sanitize
sys.modules["services"] = services_pkg

_src = _load("resume_source", "services/resume_source.py")
sys.modules["services.resume_source"] = _src
sys.modules["resume_source"] = _src

build_path = _src.build_path
detect_skill_keys_in_text = _src.detect_skill_keys_in_text
extract_text_from_docx = _src.extract_text_from_docx
extract_text_from_upload = _src.extract_text_from_upload
normalize_source_text = _src.normalize_source_text
SOURCE_EVIDENCE = _src.SOURCE_EVIDENCE


def test_python_in_text_detects_skill():
    keys = detect_skill_keys_in_text("Опыт: Python, REST API")
    assert "python" in keys


def test_word_in_text_gives_evidenced_skill_without_learn():
    path = build_path("Работал с Python и FastAPI", "backend", [])
    keys = {s["key"] for s in path["skills"]}
    assert "python" in keys
    assert "fastapi" in keys
    python = next(s for s in path["skills"] if s["key"] == "python")
    assert python["level"] == "junior"
    assert SOURCE_EVIDENCE in python["evidence"]


def test_uncompleted_role_lecture_does_not_create_skill():
    path = build_path("Меня зовут Иван, ищу работу", "backend", [])
    keys = {s["key"] for s in path["skills"]}
    # No matching words → no skills from text; uncompleted lectures do not invent skills
    assert "postgresql" not in keys
    assert "docker" not in keys
    # Recommendations still list role modules
    rec_keys = {r["skillKey"] for r in path["recommendations"]}
    assert "python" in rec_keys
    assert "postgresql" in rec_keys
    assert all(not r["evidenced"] or r["skillKey"] in keys for r in path["recommendations"])


def test_completed_slug_gives_skill_without_word_in_text():
    path = build_path("Обычный текст без технологий", "backend", ["b03-python"])
    keys = {s["key"] for s in path["skills"]}
    assert "python" in keys
    python = next(s for s in path["skills"] if s["key"] == "python")
    assert python["sourceSlugs"] == ["b03-python"]
    assert python.get("evidenceSource") == "learn"


def test_empty_text_rejected():
    try:
        normalize_source_text("   ")
        assert False, "expected ValueError"
    except ValueError as exc:
        assert "текст" in str(exc).lower() or "нужен" in str(exc).lower()


def test_empty_docx_rejected():
    from docx import Document

    buf = io.BytesIO()
    Document().save(buf)
    data = buf.getvalue()
    try:
        extract_text_from_upload(filename="empty.docx", data=data)
        assert False, "expected ValueError"
    except ValueError as exc:
        assert "текст" in str(exc).lower() or "извлечь" in str(exc).lower()


def test_docx_with_python_extracts_and_detects():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Backend developer. Stack: Python, Docker.")
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_text_from_docx(buf.getvalue())
    assert "Python" in text
    keys = detect_skill_keys_in_text(text)
    assert "python" in keys
    assert "docker" in keys


def test_unsupported_extension_rejected():
    try:
        extract_text_from_upload(filename="resume.txt", data=b"hello python")
        assert False, "expected ValueError"
    except ValueError as exc:
        assert "pdf" in str(exc).lower() or "docx" in str(exc).lower()
