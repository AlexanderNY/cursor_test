"""Build a resume file hh.ru can import (DOCX parses best; PDF is the same text).

hh.ru → Создать резюме → Загрузить файл.
Accepted: DOC, DOCX, PDF, RTF, ODT, usually up to 5 MB.
The recognizer expects a single column and the headings it already prints
on its own resumes: желаемая должность, опыт работы, образование,
ключевые навыки, обо мне. Dates look like «Январь 2020 — настоящее время».
"""
from __future__ import annotations

import html
import re
from io import BytesIO
from pathlib import Path
from typing import Any, Optional

from services.text_sanitize import sanitize_plain_text

TEMPLATES_DIR = Path(__file__).resolve().parents[1] / "templates"

MAX_EXPERIENCE = 12
MAX_EDUCATION = 6

MONTHS_NOM = (
    "",
    "январь",
    "февраль",
    "март",
    "апрель",
    "май",
    "июнь",
    "июль",
    "август",
    "сентябрь",
    "октябрь",
    "ноябрь",
    "декабрь",
)
MONTHS_GEN = (
    "",
    "января",
    "февраля",
    "марта",
    "апреля",
    "мая",
    "июня",
    "июля",
    "августа",
    "сентября",
    "октября",
    "ноября",
    "декабря",
)

# Wording hh.ru uses in the resume wizard, so the file parser can map checkboxes.
EMPLOYMENT_HH: dict[str, str] = {
    "full": "полная занятость",
    "part": "частичная занятость",
    "project": "проектная работа",
    "volunteer": "волонтерство",
    "internship": "стажировка",
}
SCHEDULE_HH: dict[str, str] = {
    "office": "полный день",
    "remote": "удаленная работа",
    "hybrid": "гибкий график",
    "travel": "вахтовый метод",
}
EDUCATION_LEVEL_HH: dict[str, str] = {
    "secondary": "Среднее",
    "special_secondary": "Среднее специальное",
    "unfinished_higher": "Неоконченное высшее",
    "higher": "Высшее",
    "bachelor": "Бакалавр",
    "master": "Магистр",
    "candidate": "Кандидат наук",
    "doctor": "Доктор наук",
}

_YEAR_MONTH = re.compile(r"^(\d{4})-(\d{2})$")
_ISO_DATE = re.compile(r"^(\d{4})-(\d{2})-(\d{2})")


def _parse_year_month(value: str) -> Optional[tuple[int, int]]:
    match = _YEAR_MONTH.match((value or "").strip())
    if not match:
        return None
    year = int(match.group(1))
    month = int(match.group(2))
    if year < 1950 or year > 2100 or month < 1 or month > 12:
        return None
    return year, month


def _year_month_label(year: int, month: int) -> str:
    return f"{MONTHS_NOM[month].capitalize()} {year}"


def format_period(start: str, end: str) -> str:
    parsed_start = _parse_year_month(start)
    if not parsed_start:
        return ""
    start_label = _year_month_label(*parsed_start)
    parsed_end = _parse_year_month(end) if end else None
    if not parsed_end:
        return f"{start_label} — настоящее время"
    return f"{start_label} — {_year_month_label(*parsed_end)}"


def format_birth(value: Any) -> str:
    text = str(value or "").strip()
    match = _ISO_DATE.match(text)
    if not match:
        return text
    year = int(match.group(1))
    month = int(match.group(2))
    day = int(match.group(3))
    if month < 1 or month > 12 or day < 1 or day > 31:
        return text
    return f"{day} {MONTHS_GEN[month]} {year}"


def format_salary(amount: Any, currency: Any) -> Optional[str]:
    if amount is None or amount == "":
        return None
    try:
        value = int(amount)
    except (TypeError, ValueError):
        return None
    if value <= 0:
        return None
    grouped = f"{value:,}".replace(",", " ")
    code = str(currency or "RUB").strip().upper()
    unit = "руб." if code in {"RUB", "RUR", "₽"} else code
    return f"{grouped} {unit}"


def _clip(value: Any, limit: int) -> str:
    return sanitize_plain_text(str(value or ""))[:limit].strip()


def normalize_experience(raw: Any) -> list[dict[str, str]]:
    if not isinstance(raw, list):
        return []
    out: list[dict[str, str]] = []
    for item in raw:
        if not isinstance(item, dict):
            continue
        company = _clip(item.get("company"), 160)
        position = _clip(item.get("position"), 160)
        city = _clip(item.get("city"), 120)
        description = _clip(item.get("description"), 4000)
        start = _parse_year_month(str(item.get("start") or ""))
        end = _parse_year_month(str(item.get("end") or ""))
        if not company or not position or not start:
            continue
        if end and (end[0], end[1]) < start:
            start, end = end, start
        out.append(
            {
                "company": company,
                "position": position,
                "city": city,
                "start": f"{start[0]:04d}-{start[1]:02d}",
                "end": f"{end[0]:04d}-{end[1]:02d}" if end else "",
                "description": description,
            }
        )
        if len(out) >= MAX_EXPERIENCE:
            break
    return out


def normalize_education(raw: Any) -> list[dict[str, Any]]:
    if not isinstance(raw, list):
        return []
    out: list[dict[str, Any]] = []
    for item in raw:
        if not isinstance(item, dict):
            continue
        institution = _clip(item.get("institution"), 200)
        if not institution:
            continue
        level_key = str(item.get("level") or "").strip()
        if level_key not in EDUCATION_LEVEL_HH:
            level_key = ""
        faculty = _clip(item.get("faculty"), 160)
        specialty = _clip(item.get("specialty"), 160)
        year: Optional[int] = None
        year_raw = item.get("year")
        if year_raw is not None and str(year_raw).strip():
            try:
                parsed_year = int(year_raw)
            except (TypeError, ValueError):
                parsed_year = 0
            if 1950 <= parsed_year <= 2100:
                year = parsed_year
        out.append(
            {
                "level": level_key,
                "institution": institution,
                "faculty": faculty,
                "specialty": specialty,
                "year": year,
            }
        )
        if len(out) >= MAX_EDUCATION:
            break
    return out


def _full_name(profile: dict[str, Any], username: str) -> str:
    parts = [
        str(profile.get("lastName") or ""),
        str(profile.get("firstName") or ""),
        str(profile.get("patronymic") or ""),
    ]
    joined = " ".join(part for part in parts if part).strip()
    return joined or username or "Кандидат"


def _skill_name(skill: dict[str, Any]) -> str:
    name = str(skill.get("name") or "").strip()
    if name:
        return name
    display = str(skill.get("display") or "").strip()
    if " — " in display:
        return display.split(" — ", 1)[0].strip()
    return display


def _selected_projects(projects: Any, *, max_count: int = 6) -> list[dict[str, Any]]:
    if not isinstance(projects, list):
        return []
    out: list[dict[str, Any]] = []
    for raw in projects:
        if not isinstance(raw, dict) or not raw.get("selected"):
            continue
        name = str(raw.get("name") or "").strip()
        url = str(raw.get("url") or "").strip()
        if not name or not url:
            continue
        out.append(
            {
                "name": name,
                "url": url,
                "description": str(raw.get("description") or "").strip(),
                "language": str(raw.get("language") or "").strip(),
                "stars": int(raw.get("stars") or 0),
            }
        )
        if len(out) >= max_count:
            break
    return out


def _labeled(values: Any, labels: dict[str, str]) -> str:
    if not isinstance(values, list):
        return ""
    words: list[str] = []
    seen: set[str] = set()
    for raw in values:
        key = str(raw or "").strip()
        label = labels.get(key, "")
        if not label or label in seen:
            continue
        seen.add(label)
        words.append(label)
    return ", ".join(words)


def _education_detail(item: dict[str, Any]) -> str:
    faculty = str(item.get("faculty") or "").strip()
    specialty = str(item.get("specialty") or "").strip()
    if faculty and specialty:
        return f"{faculty}, {specialty}"
    return faculty or specialty


def prepare_resume_view(
    *,
    profile: dict[str, Any],
    resume: dict[str, Any],
    skills: list[dict[str, Any]],
    username: str,
) -> dict[str, Any]:
    name = _full_name(profile, username)
    contacts: list[str] = []
    phone = str(profile.get("phone") or "").strip()
    email = str(profile.get("email") or "").strip()
    city = str(profile.get("city") or "").strip()
    birth = format_birth(profile.get("birthDate"))
    citizenship = str(profile.get("citizenship") or "").strip()
    if phone:
        contacts.append(f"Телефон: {phone}")
    if email:
        contacts.append(f"Электронная почта: {email}")
    if city:
        contacts.append(f"Проживает: {city}")
    if birth:
        contacts.append(f"Дата рождения: {birth}")
    if citizenship:
        contacts.append(f"Гражданство: {citizenship}")
    contacts.append(
        "Командировки: готов" if profile.get("readyForTrips") else "Командировки: не готов"
    )

    position = str(resume.get("specialization") or resume.get("title") or "").strip()
    salary = format_salary(resume.get("salaryAmount"), resume.get("salaryCurrency"))
    employment = _labeled(resume.get("employmentTypes"), EMPLOYMENT_HH)
    schedule = _labeled(resume.get("workFormats"), SCHEDULE_HH)
    about = sanitize_plain_text(str(resume.get("about") or ""))

    skill_names: list[str] = []
    seen_skills: set[str] = set()
    for skill in skills:
        if not isinstance(skill, dict):
            continue
        label = _skill_name(skill)
        key = label.casefold()
        if not label or key in seen_skills:
            continue
        seen_skills.add(key)
        skill_names.append(label)

    experience_view: list[dict[str, str]] = []
    for item in normalize_experience(resume.get("experience")):
        period = format_period(item["start"], item["end"])
        description = item["description"]
        experience_view.append(
            {
                "period": period,
                "company": item["company"],
                "city": item["city"],
                "position": item["position"],
                "description": description,
                "description_html": html.escape(description).replace("\n", "<br/>"),
            }
        )

    education_view: list[dict[str, str]] = []
    for item in normalize_education(resume.get("education")):
        level_key = str(item.get("level") or "")
        year = item.get("year")
        education_view.append(
            {
                "level": EDUCATION_LEVEL_HH.get(level_key, ""),
                "year": str(year) if year else "",
                "institution": str(item["institution"]),
                "detail": _education_detail(item),
            }
        )

    titles = resume.get("badgeTitles")
    if isinstance(titles, list) and titles:
        badges = [str(title).strip() for title in titles if str(title).strip()]
    else:
        badges = [
            str(raw).strip()
            for raw in (resume.get("selectedBadgeIds") or [])
            if str(raw).strip()
        ]

    return {
        "name": name,
        "contacts": contacts,
        "position": position,
        "salary": salary or "",
        "employment": employment,
        "schedule": schedule,
        "about": about,
        "about_html": html.escape(about).replace("\n", "<br/>") if about else "",
        "skills": skill_names,
        "skills_line": ", ".join(skill_names),
        "experience": experience_view,
        "education": education_view,
        "badges": badges,
        "projects": _selected_projects(resume.get("githubProjects") or []),
    }


def render_resume_html(
    *,
    profile: dict[str, Any],
    resume: dict[str, Any],
    skills: list[dict[str, Any]],
    username: str,
) -> str:
    from jinja2 import Environment, FileSystemLoader, select_autoescape

    view = prepare_resume_view(
        profile=profile, resume=resume, skills=skills, username=username
    )
    env = Environment(
        loader=FileSystemLoader(str(TEMPLATES_DIR)),
        autoescape=select_autoescape(["html", "xml"]),
    )
    template = env.get_template("resume_hh.html")
    return template.render(**view)


async def build_pdf_bytes(html_doc: str) -> bytes:
    from playwright.async_api import async_playwright

    async with async_playwright() as playwright:
        browser = await playwright.chromium.launch(
            headless=True,
            args=["--no-sandbox", "--disable-dev-shm-usage"],
        )
        try:
            page = await browser.new_page()
            await page.set_content(html_doc, wait_until="networkidle")
            pdf = await page.pdf(
                format="A4",
                print_background=True,
                margin={"top": "16mm", "bottom": "16mm", "left": "14mm", "right": "14mm"},
            )
            return bytes(pdf)
        finally:
            await browser.close()


def _add_heading(doc: Any, text: str) -> None:
    from docx.shared import Pt, RGBColor

    heading = doc.add_heading(text, level=1)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
        run.font.size = Pt(14)
        run.bold = True


def build_docx_bytes(
    *,
    profile: dict[str, Any],
    resume: dict[str, Any],
    skills: list[dict[str, Any]],
    username: str,
) -> bytes:
    from docx import Document
    from docx.shared import Pt

    view = prepare_resume_view(
        profile=profile, resume=resume, skills=skills, username=username
    )
    doc = Document()
    doc.core_properties.author = "9to18.ru"
    doc.core_properties.title = view["name"]
    doc.core_properties.subject = "Резюме для загрузки на hh.ru"
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    name_paragraph = doc.add_paragraph()
    name_run = name_paragraph.add_run(view["name"])
    name_run.bold = True
    name_run.font.size = Pt(18)

    for line in view["contacts"]:
        doc.add_paragraph(line)

    _add_heading(doc, "Желаемая должность и зарплата")
    doc.add_paragraph(view["position"] or "—")
    if view["salary"]:
        doc.add_paragraph(view["salary"])
    if view["employment"]:
        doc.add_paragraph(f"Занятость: {view['employment']}")
    if view["schedule"]:
        doc.add_paragraph(f"График работы: {view['schedule']}")

    if view["experience"]:
        _add_heading(doc, "Опыт работы")
        for job in view["experience"]:
            if job["period"]:
                period = doc.add_paragraph()
                period_run = period.add_run(job["period"])
                period_run.bold = True
            company = doc.add_paragraph()
            company.add_run(job["company"]).bold = True
            if job["city"]:
                doc.add_paragraph(job["city"])
            doc.add_paragraph(job["position"])
            if job["description"]:
                for block in str(job["description"]).split("\n"):
                    text = block.strip()
                    if text:
                        doc.add_paragraph(text)

    if view["education"]:
        _add_heading(doc, "Образование")
        for item in view["education"]:
            if item["level"]:
                level = doc.add_paragraph()
                level.add_run(item["level"]).bold = True
            if item["year"]:
                doc.add_paragraph(item["year"])
            school = doc.add_paragraph()
            school.add_run(item["institution"]).bold = True
            if item["detail"]:
                doc.add_paragraph(item["detail"])

    if view["skills_line"]:
        _add_heading(doc, "Ключевые навыки")
        doc.add_paragraph(view["skills_line"])

    if view["about"]:
        _add_heading(doc, "Обо мне")
        for block in str(view["about"]).split("\n"):
            text = block.strip()
            if text:
                doc.add_paragraph(text)

    if view["badges"]:
        _add_heading(doc, "Достижения")
        for label in view["badges"]:
            doc.add_paragraph(label, style="List Bullet")

    if view["projects"]:
        _add_heading(doc, "Проекты")
        for project in view["projects"]:
            line = str(project.get("name") or "")
            language = str(project.get("language") or "").strip()
            if language:
                line += f" · {language}"
            stars = int(project.get("stars") or 0)
            if stars:
                line += f" · ★{stars}"
            doc.add_paragraph(line, style="List Bullet")
            description = str(project.get("description") or "").strip()
            if description:
                doc.add_paragraph(description)
            url = str(project.get("url") or "").strip()
            if url:
                doc.add_paragraph(url)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
