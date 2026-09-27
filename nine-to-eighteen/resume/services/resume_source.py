"""Parse uploaded resume text and build evidenced skills + Learn recommendations."""
from __future__ import annotations

import io
import re
from typing import Any, Optional

from services.resume_skills import (
    GENERAL_FALLBACK_KEYS,
    SKILL_RULES,
    SkillRule,
    generate_skills_from_progress,
    missing_slugs_for_upgrade,
    role_skill_keys,
    skill_level_and_evidence,
)
from services.text_sanitize import sanitize_plain_text

MAX_SOURCE_LEN = 50_000
MAX_UPLOAD_BYTES = 8 * 1024 * 1024
SOURCE_EVIDENCE = "указано в загруженном резюме"

# Aliases matched with word boundaries (case-insensitive).
SKILL_ALIASES: dict[str, tuple[str, ...]] = {
    "python": ("python", "питон", "пайтон"),
    "fastapi": ("fastapi",),
    "react": ("react", "reactjs", "react.js"),
    "postgresql": ("postgresql", "postgres", "psql", r"\bsql\b"),
    "docker": ("docker", "dockerfile", "docker-compose", "compose"),
    "kubernetes": ("kubernetes", "k8s", "kubectl"),
    "cicd": ("ci/cd", "cicd", "github actions", "gitlab ci", "jenkins"),
    "networking": ("networking", "tcp/ip", "http", "dns", "сети"),
    "git": ("git", "github", "gitlab"),
    "api_test": ("insomnia", "postman", "api-тест", "api testing", "swagger"),
    "regex": ("regex", "regexp", "регулярк"),
    "redis": ("redis",),
    "architecture": ("архитектур", "microservices", "микросервис"),
    "security": ("jwt", "oauth", "https", "безопасность"),
    "qa": ("qa", "тестирован", "pytest", "e2e", "тест-дизайн"),
    "product": ("product owner", "product manager", "бэклог", "backlog", "системн.*аналит"),
    "soft_skills": ("soft skills", "soft-skills", "коммуникац"),
}


def extract_text_from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        try:
            page_text = page.extract_text() or ""
        except Exception:
            page_text = ""
        if page_text.strip():
            parts.append(page_text)
    return "\n".join(parts)


def extract_text_from_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [((cell.text or "").strip()) for cell in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    return "\n".join(parts)


def extract_text_from_upload(
    *,
    filename: str,
    data: bytes,
    content_type: str = "",
) -> str:
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError(f"Файл слишком большой (макс. {MAX_UPLOAD_BYTES // (1024 * 1024)} МБ)")
    if not data:
        raise ValueError("Пустой файл")

    name = (filename or "").strip().lower()
    ctype = (content_type or "").strip().lower()
    is_pdf = name.endswith(".pdf") or "pdf" in ctype
    is_docx = (
        name.endswith(".docx")
        or "wordprocessingml" in ctype
        or "officedocument" in ctype and "word" in ctype
    )
    if is_pdf:
        raw = extract_text_from_pdf(data)
    elif is_docx or name.endswith(".doc"):
        if name.endswith(".doc") and not name.endswith(".docx"):
            raise ValueError("Поддерживаются PDF и DOCX (не старый .doc)")
        raw = extract_text_from_docx(data)
    else:
        raise ValueError("Поддерживаются только PDF и DOCX")

    text = sanitize_plain_text(raw, max_length=MAX_SOURCE_LEN)
    if not text:
        raise ValueError("Не удалось извлечь текст из файла")
    return text


def normalize_source_text(text: str) -> str:
    cleaned = sanitize_plain_text(text or "", max_length=MAX_SOURCE_LEN)
    if not cleaned:
        raise ValueError("Нужен текст резюме")
    return cleaned


def _alias_patterns(aliases: tuple[str, ...]) -> list[re.Pattern[str]]:
    patterns: list[re.Pattern[str]] = []
    for alias in aliases:
        if alias.startswith(r"\b") or ".*" in alias:
            patterns.append(re.compile(alias, re.IGNORECASE))
            continue
        escaped = re.escape(alias)
        patterns.append(re.compile(rf"(?<!\w){escaped}(?!\w)", re.IGNORECASE))
    return patterns


_COMPILED_ALIASES: dict[str, list[re.Pattern[str]]] = {
    key: _alias_patterns(aliases if isinstance(aliases, tuple) else (aliases,))
    for key, aliases in SKILL_ALIASES.items()
}


def detect_skill_keys_in_text(source_text: str) -> set[str]:
    text = source_text or ""
    if not text:
        return set()
    found: set[str] = set()
    for skill_key, patterns in _COMPILED_ALIASES.items():
        for pattern in patterns:
            if pattern.search(text):
                found.add(skill_key)
                break
    return found


def _rule_by_key(key: str) -> Optional[SkillRule]:
    return next((r for r in SKILL_RULES if r.key == key), None)


def skill_from_source_text(rule: SkillRule) -> dict[str, Any]:
    return {
        "key": rule.key,
        "name": rule.name,
        "level": "junior",
        "evidence": SOURCE_EVIDENCE,
        "sourceSlugs": [],
        "mapBranch": rule.map_branch or None,
        "display": f"{rule.name} — junior ({SOURCE_EVIDENCE})",
        "evidenceSource": "resume_text",
    }


def merge_evidenced_skills(
    *,
    from_progress: list[dict[str, Any]],
    from_text_keys: set[str],
) -> list[dict[str, Any]]:
    """Prefer Learn-progress skills; add text-only skills as junior."""
    by_key: dict[str, dict[str, Any]] = {}
    for skill in from_progress:
        key = str(skill.get("key") or "")
        if not key:
            continue
        item = dict(skill)
        item["evidenceSource"] = "learn"
        by_key[key] = item

    for key in sorted(from_text_keys):
        if key in by_key:
            continue
        rule = _rule_by_key(key)
        if rule is None:
            continue
        by_key[key] = skill_from_source_text(rule)

    # Stable order: SKILL_RULES order
    order = {r.key: i for i, r in enumerate(SKILL_RULES)}
    return sorted(by_key.values(), key=lambda s: order.get(str(s.get("key")), 999))


def recommendation_keys_for_role(
    role_track: str,
    *,
    text_keys: set[str],
) -> list[str]:
    role = (role_track or "").strip().lower() or "general"
    keys = list(role_skill_keys(role))
    if role == "general" or not keys:
        keys = sorted(text_keys) if text_keys else list(GENERAL_FALLBACK_KEYS)
    # Preserve role order, append any extra text keys not in role list
    seen = set(keys)
    for key in sorted(text_keys):
        if key not in seen:
            keys.append(key)
            seen.add(key)
    return keys


def build_path(
    source_text: str,
    role_track: str,
    completed_slugs: list[str] | set[str],
) -> dict[str, Any]:
    """Evidenced skills + Learn recommendations for the guided path."""
    completed = {str(s).strip() for s in completed_slugs if str(s).strip()}
    text_keys = detect_skill_keys_in_text(source_text)
    from_progress = generate_skills_from_progress(completed)
    skills = merge_evidenced_skills(from_progress=from_progress, from_text_keys=text_keys)
    evidenced_keys = {str(s.get("key")) for s in skills}

    rec_keys = recommendation_keys_for_role(role_track, text_keys=text_keys)
    recommendations: list[dict[str, Any]] = []
    for key in rec_keys:
        rule = _rule_by_key(key)
        if rule is None:
            continue
        missing = missing_slugs_for_upgrade(key, completed)
        if not missing:
            continue
        is_evidenced = key in evidenced_keys
        if is_evidenced:
            reason = f"Усильте «{rule.name}» до middle: пройдите оставшиеся модули Learn."
        else:
            reason = (
                f"Для роли нужны модули по «{rule.name}» — "
                "после прохождения навык появится в резюме."
            )
        recommendations.append(
            {
                "skillKey": key,
                "name": rule.name,
                "learnSlugs": list(missing),
                "evidenced": is_evidenced,
                "reason": reason,
            }
        )

    return {
        "skills": skills,
        "recommendations": recommendations,
        "detectedSkillKeys": sorted(text_keys),
        "roleTrack": (role_track or "").strip().lower() or "general",
    }
